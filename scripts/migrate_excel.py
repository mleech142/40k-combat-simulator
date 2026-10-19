#!/usr/bin/env python
"""
Build the app's data JSON from the v2 Excel workbook.

Reads 40KCombatSimulator_v2.xlsx (units, loadouts, weapons, faction rules and
stratagems from each faction tab, core stratagems from the Keywords tab) and writes
data/40k_data.json. The app loads that file on every start and edits it from then on,
so this script never touches the database.

Usage:
    python scripts/migrate_excel.py                  # refuses if data/40k_data.json exists
    python scripts/migrate_excel.py --force          # rebuild it from Excel (app edits are lost)
    python scripts/migrate_excel.py --output x.json  # write elsewhere, e.g. to compare
    python scripts/migrate_excel.py --reference-only # refresh only the read-only reference data (keywords,
                                                     # generic target profiles, role benchmarks) in the existing JSON

IDs are stable: anything whose name matches an entry in the previous JSON keeps its ID,
so saved scenarios and comparisons keep pointing at the right units and weapons.
Names are written exactly as in the workbook; likely problems are listed in the report.
"""

import argparse
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'src'))

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from combat_sim.data_store import DEFAULT_DATA_PATH, SCHEMA_VERSION, read_data_file, write_data_file
from combat_sim.simulation import check_keywords

DEFAULT_WORKBOOK = ROOT / '40KCombatSimulator_v2.xlsx'

HEADER_ROW = 3
FIRST_DATA_ROW = 4

# Faction tab layout (v2). Header text per column, checked before reading.
UNIT_COLS = {
    'name': ('B', 'Name'), 'toughness': ('C', 'T'), 'wounds': ('D', 'W'), 'save': ('E', 'Sv'),
    'invulnerable_save': ('F', 'InSv'), 'feel_no_pain': ('G', 'FNP'), 'modifiers': ('H', 'Modifiers'),
    'move': ('I', 'M'), 'oc': ('J', 'OC'), 'leadership': ('K', 'Ld'), 'points_cost': ('L', 'Points'),
    'default_models': ('M', 'Models'), 'keywords': ('N', 'Keywords'), 'can_lead': ('O', 'Leader (can lead)'),
    'abilities': ('P', 'Abilities'),
}
LOADOUT_COLS = {'weapon': ('Q', 'Weapon'), 'type': ('R', 'Type'), 'qty': ('S', 'Qty')}
RANGED_COLS = {'name': ('U', 'Weapon'), 'attacks': ('V', 'A'), 'skill': ('W', 'BS'), 'strength': ('X', 'S'),
               'ap': ('Y', 'AP'), 'damage': ('Z', 'D'), 'abilities': ('AA', 'Modifier')}
MELEE_COLS = {'name': ('AD', 'Weapon'), 'attacks': ('AE', 'A'), 'skill': ('AF', 'BS'), 'strength': ('AG', 'S'),
              'ap': ('AH', 'AP'), 'damage': ('AI', 'D'), 'abilities': ('AJ', 'Modifier')}
RULE_COLS = {'name': ('AM', 'Rule'), 'applies_to': ('AN', 'Applies To'), 'phase': ('AO', 'Phase'),
             'modifier': ('AP', 'Modifier')}
STRAT_COLS = {'name': ('AR', 'Stratagem'), 'cp': ('AS', 'CP'), 'applies_to': ('AT', 'Applies To'),
              'phase': ('AU', 'Phase'), 'modifier': ('AV', 'Modifier'), 'notes': ('AW', 'Notes')}


class Report:
    def __init__(self):
        self.lines = {}

    def add(self, section, message):
        self.lines.setdefault(section, []).append(message)

    def print(self):
        for section, messages in self.lines.items():
            print(f'\n{section} ({len(messages)}):')
            for message in messages:
                print(f'  - {message}')


# ---------------------------------------------------------------------------
# Cell parsing
# ---------------------------------------------------------------------------

def clean(value):
    """Stripped text, or None for blanks and '-'."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return None if value in ('', '-') else value
    return value


def to_int(value):
    """Integer from 3, 3.0, '3', '3+' or '6"'; None for blanks and '-'."""
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    match = re.match(r'^-?\d+', str(value))
    return int(match.group()) if match else None


def dice(value):
    """Attacks/Damage as text: 2 -> '2', 'd6+1' -> 'D6+1'."""
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().upper().replace(' ', '')


def text_list(value):
    value = clean(value)
    if value is None:
        return []
    return [part.strip() for part in str(value).split(',') if part.strip() and part.strip() != '-']


def cell(ws, row, col_letter):
    return ws.cell(row, column_index_from_string(col_letter)).value


def check_headers(ws, *col_maps):
    problems = []
    for col_map in col_maps:
        for letter, expected in col_map.values():
            actual = cell(ws, HEADER_ROW, letter)
            actual = actual.strip() if isinstance(actual, str) else actual
            if actual != expected:
                problems.append(f'{letter}{HEADER_ROW} is {actual!r}, expected {expected!r}')
    if problems:
        raise SystemExit(f'Sheet "{ws.title}" does not match the v2 layout:\n  ' + '\n  '.join(problems))


# ---------------------------------------------------------------------------
# Stable IDs
# ---------------------------------------------------------------------------

class IdAllocator:
    """
    Reuses IDs from the previous JSON by key; new keys get max+1.

    A key can have several IDs (e.g. a unit name repeated for different loadouts);
    they are handed out in their original order.
    """

    def __init__(self, previous):
        self.previous = previous  # key -> [ids in ID order]
        self.next_id = max((i for ids in previous.values() for i in ids), default=0) + 1

    def get(self, key):
        if self.previous.get(key):
            return self.previous[key].pop(0)
        new_id = self.next_id
        self.next_id += 1
        return new_id


def previous_id_maps(data):
    faction_names = {f['id']: f['name'] for f in data.get('factions', [])}

    def collect(records, key):
        ids = {}
        for record in sorted(records, key=lambda r: r['id']):
            ids.setdefault(key(record), []).append(record['id'])
        return ids

    by_faction_and_name = lambda r: (faction_names.get(r['faction_id']), r['name'])
    return {
        'factions': collect(data.get('factions', []), lambda f: f['name']),
        'units': collect(data.get('units', []), by_faction_and_name),
        'ranged_weapons': collect(data.get('ranged_weapons', []), by_faction_and_name),
        'melee_weapons': collect(data.get('melee_weapons', []), by_faction_and_name),
        'faction_rules': collect(data.get('faction_rules', []), by_faction_and_name),
        'stratagems': collect(data.get('stratagems', []), by_faction_and_name),
        'keywords': collect(data.get('keywords', []), lambda k: k['name']),
        'target_profiles': collect(data.get('target_profiles', []), lambda t: (t['kind'], t['name'])),
    }


# ---------------------------------------------------------------------------
# Reading the workbook
# ---------------------------------------------------------------------------

def read_table(ws, table_name):
    """Rows of an Excel table as dicts keyed by header."""
    table = ws.tables[table_name]
    rows = list(ws[table.ref])
    headers = [str(c.value).strip() for c in rows[0]]
    return [dict(zip(headers, (c.value for c in row))) for row in rows[1:]]


def read_weapons(ws, cols, faction, weapon_type, ids, report):
    weapons = {}
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        name = clean(cell(ws, row, cols['name'][0]))
        if not name:
            continue
        name = str(name)
        if name in weapons:
            report.add('Duplicate weapon names (first kept)', f'{faction["name"]}: {weapon_type} "{name}" (row {row})')
            continue
        skill = to_int(cell(ws, row, cols['skill'][0]))
        record = {
            'id': ids.get((faction['name'], name)),
            'faction_id': faction['id'],
            'name': name,
            'attacks': dice(cell(ws, row, cols['attacks'][0])) or '1',
            'strength': to_int(cell(ws, row, cols['strength'][0])) or 0,
            'armor_penetration': abs(to_int(cell(ws, row, cols['ap'][0])) or 0),
            'damage': dice(cell(ws, row, cols['damage'][0])) or '1',
            'abilities': text_list(cell(ws, row, cols['abilities'][0])),
        }
        if weapon_type == 'ranged':
            record['ballistic_skill'] = skill
            record['range_inches'] = None  # Not in the workbook; filled in from the app
        else:
            record['weapon_skill'] = skill
        weapons[name] = record
    return weapons


def read_faction(ws, faction, ids, known_keywords, report):
    check_headers(ws, UNIT_COLS, LOADOUT_COLS, RANGED_COLS, MELEE_COLS, RULE_COLS, STRAT_COLS)

    ranged = read_weapons(ws, RANGED_COLS, faction, 'ranged', ids['ranged_weapons'], report)
    melee = read_weapons(ws, MELEE_COLS, faction, 'melee', ids['melee_weapons'], report)
    by_type = {'ranged': ranged, 'melee': melee}

    # Units, each followed by its loadout rows (Q:S) until the next unit
    unit_rows = [r for r in range(FIRST_DATA_ROW, ws.max_row + 1) if clean(cell(ws, r, 'B'))]
    units = []
    for i, row in enumerate(unit_rows):
        name = str(clean(cell(ws, row, 'B')))
        values = {field: cell(ws, row, letter) for field, (letter, _) in UNIT_COLS.items()}
        toughness, wounds, save = to_int(values['toughness']), to_int(values['wounds']), to_int(values['save'])
        if toughness is None or wounds is None or save is None:
            report.add('Units skipped (missing T, W or Sv)', f'{faction["name"]}: "{name}" (row {row})')
            continue

        abilities = clean(values['abilities'])
        move = clean(values['move'])
        unit = {
            'id': ids['units'].get((faction['name'], name)),
            'faction_id': faction['id'],
            'name': name,
            'toughness': toughness,
            'wounds': wounds,
            'save': save,
            'invulnerable_save': to_int(values['invulnerable_save']),
            'feel_no_pain': to_int(values['feel_no_pain']),
            'move': str(int(move)) if isinstance(move, float) else (str(move) if move is not None else None),
            'oc': to_int(values['oc']),
            'leadership': to_int(values['leadership']),
            'points_cost': to_int(values['points_cost']),
            'default_models': to_int(values['default_models']),
            'keywords': text_list(values['keywords']),
            'modifiers': text_list(values['modifiers']),
            'can_lead': text_list(values['can_lead']),
            'abilities': str(abilities) if abilities is not None else None,
            'points_by_size': [],
            'ranged_weapons': [],
            'melee_weapons': [],
        }

        if unit['points_cost'] is not None and unit['default_models']:
            unit['points_by_size'] = [{'models': unit['default_models'], 'points': unit['points_cost']}]

        end = unit_rows[i + 1] if i + 1 < len(unit_rows) else ws.max_row + 1
        loadout = {'ranged': {}, 'melee': {}}
        for lrow in range(row, end):
            weapon_name = clean(cell(ws, lrow, 'Q'))
            if not weapon_name:
                continue
            weapon_name = str(weapon_name)
            weapon_type = str(clean(cell(ws, lrow, 'R')) or '').lower()
            qty = to_int(cell(ws, lrow, 'S')) or 1
            if weapon_type not in by_type:
                report.add('Loadout rows skipped (no matching weapon)',
                           f'{faction["name"]} / {name}: "{weapon_name}" has type {cell(ws, lrow, "R")!r} (row {lrow})')
                continue
            weapon = by_type[weapon_type].get(weapon_name)
            if not weapon:
                other = 'melee' if weapon_type == 'ranged' else 'ranged'
                hint = f' (it is in the {other} table)' if weapon_name in by_type[other] else ''
                report.add('Loadout rows skipped (no matching weapon)',
                           f'{faction["name"]} / {name}: no {weapon_type} profile "{weapon_name}"{hint} (row {lrow})')
                continue
            if weapon['id'] in loadout[weapon_type]:
                report.add('Weapons listed twice for one unit (quantities added)',
                           f'{faction["name"]} / {name}: "{weapon_name}" (row {lrow})')
                loadout[weapon_type][weapon['id']] += qty
            else:
                loadout[weapon_type][weapon['id']] = qty

        unit['ranged_weapons'] = [{'weapon_id': wid, 'qty': q} for wid, q in sorted(loadout['ranged'].items())]
        unit['melee_weapons'] = [{'weapon_id': wid, 'qty': q} for wid, q in sorted(loadout['melee'].items())]
        units.append(unit)

    # Rules and stratagems
    rules, strats = [], []
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        rule_name = clean(cell(ws, row, 'AM'))
        if rule_name:
            rules.append({
                'id': ids['faction_rules'].get((faction['name'], str(rule_name))),
                'faction_id': faction['id'],
                'name': str(rule_name),
                'applies_to': clean(cell(ws, row, 'AN')),
                'phase': clean(cell(ws, row, 'AO')),
                'modifier': clean(cell(ws, row, 'AP')),
            })
        strat_name = clean(cell(ws, row, 'AR'))
        if strat_name:
            cp = clean(cell(ws, row, 'AS'))
            strats.append({
                'id': ids['stratagems'].get((faction['name'], str(strat_name))),
                'faction_id': faction['id'],
                'name': str(strat_name),
                'cp': str(int(cp)) if isinstance(cp, float) else (str(cp) if cp is not None else None),
                'applies_to': clean(cell(ws, row, 'AT')),
                'phase': clean(cell(ws, row, 'AU')),
                'modifier': clean(cell(ws, row, 'AV')),
                'notes': clean(cell(ws, row, 'AW')),
            })

    check_names(faction, units, ranged, melee, known_keywords, report)
    return units, list(ranged.values()), list(melee.values()), rules, strats


def check_names(faction, units, ranged, melee, known_keywords, report):
    """Report likely typos: unknown weapon keywords and 'can lead' names that match no unit."""
    for weapon in list(ranged.values()) + list(melee.values()):
        result = check_keywords(', '.join(weapon['abilities']), known_keywords)
        for entry, canonical in result['capitals']:
            report.add('Weapon keywords with different capitals (still recognised; consider fixing)',
                       f'{faction["name"]} / {weapon["name"]}: "{entry}" (Keywords tab has "{canonical}")')
        for entry in result['unknown']:
            report.add('Weapon keywords not on the Keywords tab', f'{faction["name"]} / {weapon["name"]}: "{entry}"')

    unit_names = {u['name'] for u in units}
    seen = {}
    for unit in units:
        seen[unit['name']] = seen.get(unit['name'], 0) + 1
    for name, count in seen.items():
        if count > 1:
            report.add('Unit names used more than once (kept as separate units)',
                       f'{faction["name"]}: "{name}" x{count}')
    for unit in units:
        for target in unit['can_lead']:
            if target not in unit_names:
                report.add('"Can lead" names that match no unit in the faction',
                           f'{faction["name"]} / {unit["name"]}: "{target}"')


def build_keywords(keywords_ws, ids):
    """The Keywords table: every rules keyword and modifier, how to write it and whether it's modelled."""
    keywords = []
    for row in read_table(keywords_ws, 'Keywords'):
        name = clean(row.get('Keyword'))
        if not name:
            continue
        keywords.append({
            'id': ids.get(str(name)),
            'name': str(name),
            'category': clean(row.get('Category')),
            'applies_to': clean(row.get('Applies To')),
            'format': clean(row.get('Format')),
            'summary': clean(row.get('Rule Summary')),
            'sim_step': clean(row.get('Sim Step')),
            'modelled': str(row.get('Modelled in Sim?') or '').strip().lower() == 'yes',
            'rule_ref': str(clean(row.get('Rule Ref'))) if clean(row.get('Rule Ref')) is not None else None,
        })
    return sorted(keywords, key=lambda k: k['id'])


def build_target_profiles(wb, ids):
    """Generic opponent profiles (Roster tab) and role-score benchmarks (Results tab)."""
    def profile(kind, name, models, t, w, sv, inv, fnp, keywords, weight=None, use=True):
        return {
            'id': ids.get((kind, name)),
            'kind': kind,
            'name': name,
            'models': to_int(models) or 1,
            'toughness': to_int(t),
            'wounds': to_int(w),
            'save': to_int(sv),
            'invulnerable_save': to_int(inv),
            'feel_no_pain': to_int(fnp),
            'keywords': text_list(keywords),
            'weight': float(weight) if isinstance(weight, (int, float)) else None,
            'use': bool(use),
        }

    profiles = []
    roster = wb['Roster']
    for row in range(55, 67):  # "Opponent - Generic Profiles": Use, Profile, Models, T, W, Sv, InSv, FNP, Keywords, Weight
        name = clean(cell(roster, row, 'C'))
        if name:
            values = [cell(roster, row, col) for col in 'DEFGHIJK']
            profiles.append(profile('generic', str(name), *values, use=cell(roster, row, 'B') is not False))

    results = wb['Results']
    for row in range(81, 88):  # Role scores: Role, T, W, Sv, InSv, FNP, Keywords, Models
        name = clean(cell(results, row, 'B'))
        if name:
            t, w, sv, inv, fnp, keywords, models = (cell(results, row, col) for col in 'CDEFGHI')
            profiles.append(profile('role', str(name), models, t, w, sv, inv, fnp, keywords))

    return sorted(profiles, key=lambda p: p['id'])


def build(workbook_path, previous, report):
    wb = load_workbook(workbook_path, data_only=True)
    keywords_ws = wb['Keywords']
    faction_names = [str(r['Faction']).strip() for r in read_table(keywords_ws, 'Factions') if clean(r['Faction'])]
    known_keywords = {str(r['Keyword']).strip() for r in read_table(keywords_ws, 'Keywords') if clean(r['Keyword'])}

    ids = {name: IdAllocator(id_map) for name, id_map in previous_id_maps(previous).items()}
    data = {
        'schema_version': SCHEMA_VERSION,
        'source': {
            'workbook': Path(workbook_path).name,
            'generated': datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        },
        'factions': [], 'units': [], 'ranged_weapons': [], 'melee_weapons': [],
        'faction_rules': [], 'stratagems': [], 'keywords': [], 'target_profiles': [],
    }

    for name in faction_names:
        if name not in wb.sheetnames:
            report.add('Factions without a tab', name)
            continue
        faction = {'id': ids['factions'].get(name), 'name': name}
        data['factions'].append(faction)
        units, ranged, melee, rules, strats = read_faction(wb[name], faction, ids, known_keywords, report)
        data['units'] += units
        data['ranged_weapons'] += ranged
        data['melee_weapons'] += melee
        data['faction_rules'] += rules
        data['stratagems'] += strats
        print(f'  {name:20} units={len(units):3}  ranged={len(ranged):3}  melee={len(melee):3}  '
              f'rules={len(rules):2}  stratagems={len(strats):2}')

    for row in read_table(keywords_ws, 'CoreStratagems'):
        strat_name = clean(row.get('Stratagem'))
        if not strat_name:
            continue
        modifier = clean(row.get('Modifier'))
        cp = clean(row.get('CP'))
        data['stratagems'].append({
            'id': ids['stratagems'].get((None, str(strat_name))),
            'faction_id': None,
            'name': str(strat_name),
            'cp': str(int(cp)) if isinstance(cp, float) else (str(cp) if cp is not None else None),
            'applies_to': clean(row.get('Applies To')),
            'phase': clean(row.get('Phase')),
            'modifier': None if modifier == 'None' else modifier,
            'notes': clean(row.get('Notes')),
        })
    print(f'  {"Core stratagems":20} {sum(1 for s in data["stratagems"] if s["faction_id"] is None)}')

    data['keywords'] = build_keywords(keywords_ws, ids['keywords'])
    print(f'  {"Keywords reference":20} {len(data["keywords"])}')
    data['target_profiles'] = build_target_profiles(wb, ids['target_profiles'])
    print(f'  {"Target profiles":20} {len(data["target_profiles"])}')

    for section in ('factions', 'units', 'ranged_weapons', 'melee_weapons', 'faction_rules', 'stratagems'):
        data[section].sort(key=lambda record: record['id'])
    return data


def main():
    parser = argparse.ArgumentParser(description='Build the data JSON from the v2 Excel workbook.')
    parser.add_argument('--workbook', default=str(DEFAULT_WORKBOOK), help='Excel workbook to read')
    parser.add_argument('--output', default=str(DEFAULT_DATA_PATH), help='JSON file to write')
    parser.add_argument('--force', action='store_true', help='overwrite the output file if it exists')
    parser.add_argument('--reference-only', action='store_true',
                        help='only refresh the read-only reference data (keywords, generic target profiles, '
                             'role benchmarks) in the existing output file; keeps all other data')
    args = parser.parse_args()

    output = Path(args.output)
    if args.reference_only:
        if not output.exists():
            print(f'{output} does not exist; run without --reference-only first.')
            return 1
        data = read_data_file(output)
        id_maps = previous_id_maps(data)
        wb = load_workbook(args.workbook, data_only=True)
        data['keywords'] = build_keywords(wb['Keywords'], IdAllocator(id_maps['keywords']))
        data['target_profiles'] = build_target_profiles(wb, IdAllocator(id_maps['target_profiles']))
        write_data_file(output, data)
        print(f'Updated the reference data in {output}: {len(data["keywords"])} keywords, '
              f'{len(data["target_profiles"])} target profiles. Nothing else was changed.')
        return 0

    if output.exists() and not args.force:
        print(f'{output} already exists. The app edits this file, so rebuilding it from Excel would lose those edits.')
        print('Use --force to overwrite it, or --output <file> to write somewhere else.')
        return 1

    # Keep IDs stable against the file being replaced, or the app's current data
    id_source = output if output.exists() else Path(DEFAULT_DATA_PATH)
    previous = read_data_file(id_source) if id_source.exists() else {}

    print(f'Reading {args.workbook}')
    report = Report()
    data = build(args.workbook, previous, report)
    write_data_file(output, data)

    print(f'\nWrote {output}')
    print(f'  {len(data["factions"])} factions, {len(data["units"])} units, '
          f'{len(data["ranged_weapons"])} ranged + {len(data["melee_weapons"])} melee weapons, '
          f'{len(data["faction_rules"])} faction rules, {len(data["stratagems"])} stratagems')
    report.print()
    return 0


if __name__ == '__main__':
    sys.exit(main())
